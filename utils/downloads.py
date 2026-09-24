"""Build a Word document from a student's responses so they can download it.

Nothing is stored on the server: the file is generated in memory and handed
straight to the student's browser.
"""

import io
from datetime import date

import streamlit as st
from docx import Document
from docx.shared import Pt

from utils.style import callout

MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

TYPING_NOTE = (
    "**Downloading takes two clicks, in this order.** First click outside the box you "
    "just typed in, or press Ctrl+Enter (Cmd+Enter on a Mac), because your text doesn't "
    "reach the app until you do. Check the counter below: it should include the box you "
    "just finished. Then click Step 1, and then Step 2."
)

PAGE_NOTE = (
    "Nothing you type here is saved automatically, and it disappears if you close the tab "
    "or reload the page. Use the **Save your work** section at the bottom to download your "
    "answers as a Word file, and download again whenever you add more."
)


def build_docx(title: str, student_name: str, sections: list[tuple[str, str]]) -> bytes:
    """sections is a list of (heading, body) pairs. Empty bodies are marked as such."""
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Georgia"
    style.font.size = Pt(11)

    doc.add_heading(title, level=1)
    meta = doc.add_paragraph()
    meta.add_run(f"Name: {student_name or '(not entered)'}\n").bold = True
    meta.add_run(f"Date: {date.today().strftime('%B %d, %Y')}")

    for heading, body in sections:
        doc.add_heading(heading, level=2)
        doc.add_paragraph(body.strip() if body and body.strip() else "(no response)")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def save_section(*, key_prefix: str, name_label: str, field_keys: list[str],
                 build, file_name: str, extra_note: str | None = None):
    """The 'Save your work' section used at the bottom of every activity page.

    The Word file is built in two steps on purpose. Streamlit prepares a download
    button's file when the page renders, so a file built in the same instant a
    student is still typing can miss their last box. Asking them to prepare the file
    first guarantees that the app has their text before the file is built.
    """
    st.header("Save your work")
    st.markdown(TYPING_NOTE)
    if extra_note:
        st.caption(extra_note)

    name = st.text_input(name_label, key=f"{key_prefix}_name")

    filled = sum(1 for k in field_keys if str(st.session_state.get(k) or "").strip())
    total = len(field_keys)
    st.progress(filled / total if total else 0.0,
                text=f"{filled} of {total} boxes on this page have text in them right now")

    # Fingerprint of everything on the page, so the app can tell when a student has
    # written more since the file was prepared.
    current = str([st.session_state.get(k) for k in field_keys] + [name])
    prepared_key = f"{key_prefix}_prepared"
    if st.button("Step 1: Prepare my Word file", key=f"{key_prefix}_prepare"):
        st.session_state[prepared_key] = current

    if st.session_state.get(prepared_key) not in (None, current):
        callout(
            "note",
            "You've changed something on the page since you prepared your file. Click "
            "**Step 1** again so the new text goes in.",
            announce=True,
        )
    elif st.session_state.get(prepared_key):
        st.download_button(
            "Step 2: Download the Word file",
            data=build(name),
            file_name=file_name,
            mime=MIME_DOCX,
            icon=":material/download:",
            type="primary",
            key=f"{key_prefix}_download",
        )
        st.caption(
            "This file holds everything the counter above is showing. If you write more "
            "afterward, click Prepare again before downloading, so the new text is included."
        )
