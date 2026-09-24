"""Renders a unit activity page from a content dictionary (see content/unit1.py etc.)
and builds the matching Word download from whatever the student typed and checked.

Block types a unit's "parts" can contain:
  steps    - numbered instructions (no input)
  refusal  - the "prefer not to use AI?" box
  sample   - the previously-generated example, in a collapsible panel
  hint     - a discussion checkpoint box
  prompt   - a question with a text box
  audit    - a question with a checkbox and a text box
  check    - a checklist item (checkbox only)
  star     - a STAR story card with several labeled text boxes
"""

import io
from datetime import date

import streamlit as st
from docx import Document
from docx.shared import Pt

from utils.downloads import PAGE_NOTE, save_section
from utils.style import banner, callout, label


def _k(unit, block_id, field=None):
    return f"{unit['key']}_{block_id}" + (f"_{field}" if field else "")


# ── Rendering ────────────────────────────────────────────────────────────────

def _render_block(unit, b):
    t = b["type"]
    if t == "steps":
        for n, item in enumerate(b["items"], start=1):
            with st.container(border=True):
                st.markdown(f"**{n}.** {item}")
    elif t == "refusal":
        callout("refusal", b["text"],
                title="Prefer not to use AI? That's always an option in this course.")
    elif t == "sample":
        with st.expander(b["label"], icon=":material/description:"):
            for sec in b["sections"]:
                if sec.get("label"):
                    label(sec["label"])
                with st.container(border=True):
                    st.markdown(sec["text"])
            st.caption(b["note"])
    elif t == "hint":
        callout("note", b["text"], title="Discussion checkpoint")
    elif t in ("prompt", "audit", "check"):
        with st.container(border=True):
            if t == "prompt":
                st.markdown(f"**{b['title']}**")
            else:
                st.checkbox(f"**{b['title']}**", key=_k(unit, b["id"], "done"))
            st.markdown(b["text"])
            if t != "check":
                st.text_area(b["title"], key=_k(unit, b["id"]), placeholder=b.get("placeholder", ""),
                             height=b.get("height", 100), label_visibility="collapsed")
    elif t == "star":
        with st.container(border=True):
            st.subheader(b["title"])
            for f in b["fields"]:
                label(f["label"])
                st.text_area(f["label"], key=_k(unit, b["id"], f["id"]),
                             placeholder=f.get("placeholder", ""), height=80,
                             label_visibility="collapsed")
                if f.get("hint"):
                    st.caption(f["hint"])


def writing_keys(unit):
    """Every text box on the page, in order, for the 'boxes filled' counter."""
    keys = []
    for part in unit["parts"]:
        for b in part["blocks"]:
            if b["type"] in ("prompt", "audit"):
                keys.append(_k(unit, b["id"]))
            elif b["type"] == "star":
                keys += [_k(unit, b["id"], f["id"]) for f in b["fields"]]
    return keys


def render_unit(unit):
    banner(unit["kicker"], unit["title"], unit["sub"])
    callout("info", PAGE_NOTE)

    with st.container(border=True):
        st.markdown(unit["goals_intro"])
        st.markdown("By the end " + unit["goals_lead"] + ":")
        st.markdown("\n".join(f"- {g}" for g in unit["goals"]))

    for part in unit["parts"]:
        st.header(f"{part['num']} · {part['title']}")
        if part.get("note"):
            st.caption(part["note"])
        for b in part["blocks"]:
            _render_block(unit, b)

    with st.container(border=True):
        label(unit["deliver_label"])
        st.markdown("\n".join(f"- {d}" for d in unit["deliver"]))

    save_section(
        key_prefix=unit["key"],
        name_label=unit["name_label"],
        field_keys=writing_keys(unit),
        build=lambda names: build_unit_docx(unit, names),
        file_name=unit["file_name"],
        extra_note="Your checklist boxes go into the file as well, marked checked or unchecked.",
    )

    st.divider()
    st.caption(unit["footer"])


# ── Word export ──────────────────────────────────────────────────────────────

def _plain(md: str) -> str:
    """Strip the light markdown used in content files so it reads cleanly in Word."""
    return md.replace("**", "").replace("*", "")


def _answer(doc, text):
    text = (text or "").strip()
    p = doc.add_paragraph(text if text else "(no response)")
    if not text:
        p.runs[0].italic = True


def build_unit_docx(unit, names: str) -> bytes:
    ss = st.session_state
    doc = Document()
    doc.styles["Normal"].font.name = "Georgia"
    doc.styles["Normal"].font.size = Pt(11)

    doc.add_heading(_plain(unit["doc_title"]), level=1)
    meta = doc.add_paragraph()
    meta.add_run(f"{unit['name_label'].split(' (')[0]}: {names or '(not entered)'}\n").bold = True
    meta.add_run(f"Date: {date.today().strftime('%B %d, %Y')}")

    for part in unit["parts"]:
        inputs = [b for b in part["blocks"] if b["type"] in ("prompt", "audit", "check", "star")]
        if not inputs:
            continue
        doc.add_heading(f"{part['num']}: {part['title']}", level=2)
        for b in inputs:
            t = b["type"]
            if t == "check":
                mark = "☑" if ss.get(_k(unit, b["id"], "done")) else "☐"
                doc.add_paragraph(f"{mark}  {_plain(b['title'])}")
                continue
            if t == "star":
                doc.add_heading(b["title"], level=3)
                for f in b["fields"]:
                    p = doc.add_paragraph()
                    p.add_run(f["label"]).bold = True
                    _answer(doc, ss.get(_k(unit, b["id"], f["id"])))
                continue
            title = _plain(b["title"])
            if t == "audit":
                title = ("☑ " if ss.get(_k(unit, b["id"], "done")) else "☐ ") + title
            doc.add_heading(title, level=3)
            q = doc.add_paragraph(_plain(b["text"]))
            q.runs[0].italic = True
            _answer(doc, ss.get(_k(unit, b["id"])))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
